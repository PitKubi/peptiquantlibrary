from django.core.management.base import BaseCommand

from catalog.models import FastaAnnotation, PublicPeptide, SourcePeptideLot
from catalog.publish import STANDARD_SEQUENCE


class Command(BaseCommand):
    help = ('Write the curation gap report for Maria as xlsx: published '
            'peptides still missing annotation, ambiguous FASTA matches, and '
            'eligible lots held back for non-standard sequences.')

    def add_arguments(self, parser):
        parser.add_argument('path', help='Output .xlsx path')

    def handle(self, *args, **options):
        import openpyxl
        workbook = openpyxl.Workbook()

        sheet = workbook.active
        sheet.title = 'Missing annotation'
        sheet.append(['Peptide sequence', 'Missing', 'SIS', 'NAT', 'Lots'])
        for p in PublicPeptide.objects.all():
            missing = [name for name, value in
                       [('gene', p.gene_name), ('uniprot', p.uniprot_id),
                        ('species', p.species)] if not value]
            if missing:
                sheet.append([p.peptide_sequence, ', '.join(missing),
                              'yes' if p.sis_available else '',
                              'yes' if p.nat_available else '', p.lot_count])

        published = set(PublicPeptide.objects.values_list(
            'peptide_sequence', flat=True))
        sheet = workbook.create_sheet('Ambiguous FASTA matches')
        sheet.append(['Peptide sequence', 'Matches in proteome(s)', 'FASTA'])
        for a in FastaAnnotation.objects.filter(match_count__gt=1):
            if a.peptide_sequence in published:
                sheet.append([a.peptide_sequence, a.match_count, a.fasta_label])

        sheet = workbook.create_sheet('Non-standard sequences')
        sheet.append(['Source key', 'Raw sequence', 'Label type', 'Status'])
        for lot in SourcePeptideLot.objects.all():
            if (lot.is_catalog_eligible()
                    and not STANDARD_SEQUENCE.match(
                        lot.peptide_sequence.strip().upper())):
                sheet.append([lot.source_key, lot.peptide_sequence,
                              lot.label_type, lot.workflow_status])

        workbook.save(options['path'])
        self.stdout.write(f'Wrote curation report to {options["path"]}')
